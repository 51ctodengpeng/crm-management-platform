# -*- coding: utf-8 -*-
"""Office 文件文本预览：不引入 python-docx / python-pptx 等重依赖

Office 2007+ 的 docx / pptx / xlsx 本质是 zip + XML：
  - docx：word/document.xml 里每个 </w:p> 是一段
  - pptx：ppt/slides/slideN.xml 里每个 <a:t> 是一段文字
  - xlsx：直接用 openpyxl（已经是导出功能的依赖）读单元格
只抽纯文本用于「快速看一眼」，排版、图片、公式一律不还原。
"""
import re
import zipfile
from html import unescape
from pathlib import Path

MAX_CHARS = 20000          # 预览文本上限，防止超大文档撑爆页面
MAX_ROWS = 200             # Excel 预览行数
MAX_COLS = 20

_TAG_RE = re.compile(r"<[^>]+>")


def _xml_text(xml: str, para_tag: str) -> str:
    """把 XML 里的段落标签换成换行，再去掉所有标签"""
    xml = re.sub(rf"</{para_tag}>", "\n", xml)
    xml = re.sub(r"<w:tab[^>]*/>", "\t", xml)
    xml = re.sub(r"<w:br[^>]*/>", "\n", xml)
    text = _TAG_RE.sub("", xml)
    text = unescape(text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _docx(path: Path) -> str:
    with zipfile.ZipFile(path) as zf:
        names = [n for n in zf.namelist() if n.startswith("word/document")]
        if not names:
            return ""
        xml = zf.read(names[0]).decode("utf-8", "ignore")
    return _xml_text(xml, "w:p")


def _pptx(path: Path) -> str:
    with zipfile.ZipFile(path) as zf:
        slides = sorted(
            (n for n in zf.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n)),
            key=lambda n: int(re.findall(r"\d+", n)[-1]),
        )
        chunks = []
        for idx, name in enumerate(slides, start=1):
            xml = zf.read(name).decode("utf-8", "ignore")
            texts = re.findall(r"<a:t>(.*?)</a:t>", xml, flags=re.S)
            body = "\n".join(unescape(_TAG_RE.sub("", t)).strip() for t in texts if t.strip())
            chunks.append(f"—— 第 {idx} 页 ——\n{body}")
    return "\n\n".join(chunks).strip()


def _xlsx(path: Path) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    chunks = []
    for ws in wb.worksheets[:3]:
        lines = [f"—— 工作表：{ws.title} ——"]
        for row_idx, row in enumerate(ws.iter_rows(max_col=MAX_COLS, values_only=True)):
            if row_idx >= MAX_ROWS:
                lines.append("…（仅预览前 200 行）")
                break
            cells = ["" if v is None else str(v) for v in row]
            if any(c.strip() for c in cells):
                lines.append(" | ".join(cells).rstrip(" |"))
        chunks.append("\n".join(lines))
    wb.close()
    return "\n\n".join(chunks).strip()


def _plain(path: Path) -> str:
    raw = path.read_bytes()[: MAX_CHARS * 4]
    for encoding in ("utf-8", "gb18030", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return ""


def extract_text(path: Path, ext: str) -> tuple[str, str]:
    """返回 (文本, 提示)；不支持的类型返回空文本"""
    ext = (ext or "").lower()
    try:
        if ext == "docx":
            text = _docx(path)
        elif ext == "pptx":
            text = _pptx(path)
        elif ext == "xlsx":
            text = _xlsx(path)
        elif ext in {"txt", "csv", "md", "log", "json"}:
            text = _plain(path)
        else:
            return "", "该类型暂不支持文本预览（doc/xls/ppt 老格式请另存为新格式）"
    except (zipfile.BadZipFile, OSError, ValueError, KeyError) as exc:
        return "", f"解析失败：{exc}"

    if not text:
        return "", "没有提取到可读文字（可能是纯图片或加密文档）"
    note = ""
    if len(text) > MAX_CHARS:
        text = text[:MAX_CHARS]
        note = f"内容较长，仅显示前 {MAX_CHARS} 字"
    return text, note
