# -*- coding: utf-8 -*-
"""路由层：数据导出模块（客户台账 / 项目台账 -> Excel）"""
from datetime import datetime
from io import BytesIO

from flask import Blueprint, request, send_file
from flask_login import login_required
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from extensions import db
from models import Customer, Payment, Project
from permissions import PERM_VIEW, apply_customer_scope, apply_project_scope
from utils import log_action

bp = Blueprint("export", __name__)

HEADER_FILL = PatternFill("solid", fgColor="006EFF")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
THIN = Side(style="thin", color="D3DCE8")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# 指定列的宽度；未指定的列按内容长度自动计算，避免新增字段后与表头错位
WIDTH_OVERRIDES = {
    "客户公司": 30, "项目名称": 30, "邮箱": 26, "公司地址": 34,
    "详细地址": 34, "客户备注": 36, "项目描述": 46, "微信号": 16,
    "所在地区": 16, "来源渠道": 14, "对接人": 12,
}


def _visual_width(text) -> int:
    return sum(2 if ord(ch) > 0x2E7F else 1 for ch in str(text))


def _auto_width(header: str, rows: list[dict]) -> int:
    if header in WIDTH_OVERRIDES:
        return WIDTH_OVERRIDES[header]
    longest = _visual_width(header)
    for row in rows:
        longest = max(longest, _visual_width(row.get(header, "")))
    # 中文按 2 个字符宽度换算，限制在 8~40 之间
    return max(8, min(40, int(longest * 1.05) + 3))


def _build_sheet(title: str, headers: list[str], rows: list[dict]) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = title

    ws.append(headers)
    for idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(idx)].width = _auto_width(header, rows)
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22

    for row in rows:
        ws.append([row.get(h, "") for h in headers])

    for r in range(2, ws.max_row + 1):
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center")
    if ws.max_row > 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _filtered_customers():
    """按列表页同样的条件导出：关键字 / 等级 / 行业 / 渠道"""
    from sqlalchemy import or_

    stmt = apply_customer_scope(db.select(Customer))
    keyword = (request.args.get("q") or "").strip()
    level = (request.args.get("level") or "").strip()
    business = (request.args.get("business") or "").strip()
    channel = (request.args.get("channel") or "").strip()
    if keyword:
        like = f"%{keyword}%"
        stmt = stmt.where(or_(
            Customer.company_name.like(like), Customer.contact_name.like(like),
            Customer.contact_phone.like(like), Customer.wechat.like(like),
            Customer.contact_email.like(like), Customer.region.like(like),
            Customer.address.like(like), Customer.remark.like(like),
            Customer.business_type.like(like), Customer.source_channel.like(like),
        ))
    if level:
        stmt = stmt.where(Customer.customer_level == level)
    if business:
        stmt = stmt.where(Customer.business_type == business)
    if channel:
        stmt = stmt.where(Customer.source_channel == channel)
    return db.session.scalars(stmt.order_by(Customer.id)).all()


def _filtered_projects():
    """按列表页同样的条件导出：关键字 / 类型 / 状态 / 服务方式 / 所属客户"""
    from sqlalchemy import or_

    stmt = db.select(Project).join(Customer, Project.customer_id == Customer.id)
    keyword = (request.args.get("q") or "").strip()
    ptype = (request.args.get("type") or "").strip()
    status = (request.args.get("status") or "").strip()
    service = (request.args.get("service") or "").strip()
    customer_id = request.args.get("customer_id", type=int)
    if keyword:
        like = f"%{keyword}%"
        stmt = stmt.where(or_(
            Project.project_name.like(like), Customer.company_name.like(like),
            Project.project_summary.like(like), Project.project_type.like(like),
            Project.project_status.like(like),
        ))
    if ptype:
        stmt = stmt.where(Project.project_type == ptype)
    if status:
        stmt = stmt.where(Project.project_status == status)
    if service:
        stmt = stmt.where(Project.service_mode == service)
    if customer_id:
        stmt = stmt.where(Project.customer_id == customer_id)
    stmt = apply_project_scope(stmt, PERM_VIEW)
    return db.session.scalars(stmt.order_by(Project.id)).all()


def _filtered_payments():
    """收款流水导出：状态 / 付款方式 / 到账日期区间 / 关键字"""
    from sqlalchemy import or_
    from sqlalchemy.orm import joinedload

    stmt = db.select(Payment).join(Project, Payment.project_id == Project.id)
    status = (request.args.get("status") or "").strip()
    method = (request.args.get("method") or "").strip()
    keyword = (request.args.get("q") or "").strip()
    date_from = request.args.get("date_from") or ""
    date_to = request.args.get("date_to") or ""
    if status:
        stmt = stmt.where(Payment.status == status)
    if method:
        stmt = stmt.where(Payment.pay_method == method)
    if keyword:
        like = f"%{keyword}%"
        stmt = stmt.where(or_(
            Project.project_name.like(like), Payment.trade_no.like(like),
            Payment.payer.like(like), Payment.remark.like(like),
        ))
    if date_from:
        stmt = stmt.where(Payment.paid_at >= date_from)
    if date_to:
        stmt = stmt.where(Payment.paid_at <= date_to)
    stmt = apply_project_scope(stmt, PERM_VIEW)
    return db.session.scalars(
        stmt.options(joinedload(Payment.project).joinedload(Project.customer))
        .order_by(Payment.paid_at, Payment.id)
    ).all()


def _stamp() -> str:
    return datetime.now().strftime("%Y%m%d_%H%M")


@bp.route("/customers.xlsx")
@login_required
def export_customers():
    customers = _filtered_customers()
    rows = [c.to_dict() for c in customers]
    # 无数据时仍导出带表头的空表，便于直接当模板使用
    headers = list(rows[0].keys()) if rows else list(Customer().to_dict().keys())
    buf = _build_sheet("客户台账", headers, rows)
    suffix = "（按当前筛选）" if request.query_string else ""
    log_action("导出", "客户", None, f"导出客户台账 {len(rows)} 条{suffix}")
    return send_file(
        buf,
        as_attachment=True,
        download_name=f"客户台账_{_stamp()}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@bp.route("/projects.xlsx")
@login_required
def export_projects():
    projects = _filtered_projects()
    rows = [p.to_dict() for p in projects]
    # 无数据时仍导出带表头的空表（用空实例取字段名，字段变更自动同步）
    headers = list(rows[0].keys()) if rows else list(Project().to_dict().keys())
    buf = _build_sheet("项目台账", headers, rows)
    suffix = "（按当前筛选）" if request.query_string else ""
    log_action("导出", "项目", None, f"导出项目台账 {len(rows)} 条{suffix}")
    return send_file(
        buf,
        as_attachment=True,
        download_name=f"项目台账_{_stamp()}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@bp.route("/payments.xlsx")
@login_required
def export_payments():
    """收款流水导出：含对账状态与经办信息，便于财务归档"""
    payments = _filtered_payments()
    rows = [p.to_dict() for p in payments]
    headers = list(rows[0].keys()) if rows else list(Payment().to_dict().keys())
    buf = _build_sheet("收款流水", headers, rows)
    suffix = "（按当前筛选）" if request.query_string else ""
    log_action("导出", "收款", None, f"导出收款流水 {len(rows)} 条{suffix}")
    return send_file(
        buf,
        as_attachment=True,
        download_name=f"收款流水_{_stamp()}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
